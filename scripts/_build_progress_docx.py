"""Generate progress_update.docx from the plain-text progress note."""
from docx import Document
from docx.shared import Pt

doc = Document()

style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)

doc.add_heading("Data-quality LLM — progress update", level=0)

doc.add_heading("TL;DR", level=1)
doc.add_paragraph(
    "End-to-end harness for LLM-based CSV error detection is up: 9-type "
    "taxonomy, per-type prompting, deterministic + LLM clustering, "
    "relation-discovery pre-pass, multi-pass Good-Turing/Chao1 false-negative "
    "guarantee, FastAPI labeling UI, and a synthetic benchmark for "
    "precision/recall/F1."
)

doc.add_heading("Taxonomy (9 types)", level=1)
doc.add_paragraph(
    "Single source of truth: error_type.csv ↔ prompts/error_types.yaml."
)
for item in [
    "Single-Column (4): typo, single_column_inconsistency (covers text / "
    "casing / pattern / punctuation / number-format), semantic_incompatible, "
    "numeric_outlier",
    "Multi-Column intra-row (3): fd_violation, row_calculation, "
    "string_calculation",
    "Multi-Row (1): row_duplication",
    "Multi-Table (1): inter_table_inconsistency",
]:
    doc.add_paragraph(item, style="List Bullet")
doc.add_paragraph(
    "One narrow LLM call per (chunk × error type), with shared global rules "
    "(R1–R7: empty default, slice-size guard, dominant ≥60% / minority ≤20% "
    "baseline, majority guard, lane discipline). Thresholds centralised so "
    "prompt edits stay in one place."
)

doc.add_heading("Pipeline", level=1)
for i, item in enumerate([
    "Per-type detection — parallel LLM calls per chunk × type.",
    "Phase A — deterministic clustering by row + overlapping columns.",
    "Phase B — LLM synthesis: pick primary_error_type, canonical "
    "explanation/fix, confidence, FP flag.",
    "Phase D (relation discovery) — one-shot per (source, error_type) call "
    "enumerates FDs / calculations / aggregations from a whole-table digest; "
    "relations are then validated row-by-row.",
    "Outputs: self-contained report.html plus clusters.jsonl / issues.jsonl "
    "/ summary.json / raw_responses/ per run.",
], start=1):
    doc.add_paragraph(item, style="List Number")
doc.add_paragraph(
    "Complementary stack: heuristic detectors (whitespace, casing, "
    "pattern_minority, FD, numeric outlier, …) and a synthetic-injection "
    "benchmark for exact per-type P/R/F1."
)

doc.add_heading("Statistical FN guarantee", level=1)
doc.add_paragraph(
    "Multi-pass with perturbed row samples; species = (source, column, row). "
    "error_type is intentionally excluded so re-labels still count as one "
    "defect."
)
for item in [
    "Good–Turing missing mass: p_miss = n1 / N (chance the next pass finds "
    "something new).",
    "Chao1 lower bound on total species → FN-rate lower bound, with "
    "Chao–Yang log-normal 95% CI as the headline guarantee.",
    "Adaptive stopping per arm = (source, error_type): stop when "
    "passes ≥ min_passes AND (GT mass < τ) and/or "
    "(Chao1 upper-CI / S_obs < ε). Loop ends when all arms saturate or hit "
    "max_passes.",
]:
    doc.add_paragraph(item, style="List Bullet")
doc.add_paragraph("The same math drives the relation-discovery loop.")

doc.add_heading("Labeling system", level=1)
doc.add_paragraph(
    "Backend: FastAPI (src/labeling/app.py). One server per project root; "
    "manages run dirs, per-run labels.jsonl, and supervised child jobs "
    "(detection, report build, relation discovery). Endpoints for runs / "
    "clusters / labels / stats / benchmark export, plus a launcher API."
)
doc.add_paragraph("Frontend (vanilla HTML/JS, CSP-locked):")
for item in [
    "Runs dashboard with tabs: Runs / Launch / Discover relations / "
    "Compare / Data prep / Jobs / Pipeline.",
    "Labeling view: filter → queue → CSV viewer with flagged cells "
    "highlighted → label as tp/fp/partial/unsure/skip with adopt-LLM-"
    "explanation or freeform note.",
    "Each report.html is self-contained with Clusters / Raw detections / "
    "Source data views, an inline localStorage review panel, and a "
    "run-parameters provenance card.",
]:
    doc.add_paragraph(item, style="List Bullet")

doc.add_heading("Current numbers (synth benchmark, 49 injections)", level=1)
for item in [
    "LLM cell-recall 0.60, heuristics 0.68, union 0.79 → complementary, "
    "not redundant.",
    "LLM wins on numeric_outlier / number_format / typo; heuristics win on "
    "whitespace.",
    "Main LLM failure mode is type-confusion, not blindness → 0.60 is a "
    "lower bound on true detection rate.",
    "Known joint blind spots: missing_value, pattern_inconsistency "
    "(fixes queued).",
]:
    doc.add_paragraph(item, style="List Bullet")

doc.add_heading("Next", level=1)
for item in [
    "Re-run heuristic-as-oracle on a less curated slice for a non-degenerate "
    "FN floor.",
    "Add missing_value to active types plus a 1-line heuristic; relax the "
    "pattern_minority threshold.",
    "Scale synth benchmark to 200+ injections / 5 spreadsheets for tighter "
    "per-type confidence intervals.",
]:
    doc.add_paragraph(item, style="List Bullet")

out = "docs/progress_update.docx"
doc.save(out)
print(f"wrote {out}")
